import docx
import os
import sys
import json
from docx2pdf import convert

# Function to replace placeholders with data and apply formatting
def replace_text_and_format(run, text, font_size=12, bold=False, italic=False):
    run.text = text  # Replace the text
    run.font.size = docx.shared.Pt(font_size)  # Set font size
    run.bold = bold  # Set bold
    run.italic = italic  # Set italic

# Function to replace placeholders in a paragraph
def replace_placeholders(paragraph, data):
    for key, value_info in data.items():
        if key in paragraph.text:
            for run in paragraph.runs:
                if key in run.text:
                    replace_text_and_format(
                        run,
                        value_info['text'],
                        font_size=value_info['font_size'],
                        bold=value_info['bold'],
                    )
def replace_placeholders_in_text(text, data):
    for key, value_info in data.items():
        text = text.replace(key, value_info['text'])  # Replace the placeholder with its value
    return text

#   # Function to convert Word document to PDF
def convert_to_pdf(word_file_path, pdf_file_path):
    convert(word_file_path, pdf_file_path)

if __name__ == "__main__":
    # Read JSON input from stdin
    input_data = json.loads(sys.stdin.read())

    # Map input data to the placeholders in the template
    data = {
        'DAY': {'text': input_data.get("agreement-day", ""), 'font_size': 12, 'bold': True},
        'MONTH': {'text': input_data.get("agreement-month", ""), 'font_size': 12, 'bold': True},
        'YEAR': {'text': input_data.get("agreement-year", ""), 'font_size': 12, 'bold': True},
        'RELEASERNAME': {'text': input_data.get("releaser-name", ""), 'font_size': 12, 'bold': True},
        'RELEASERFATHER': {'text': input_data.get("releasers-fathers-name", ""), 'font_size': 12, 'bold': True},
        'RELEASERADDRESS': {'text': input_data.get("releasers-address", ""), 'font_size': 12, 'bold': True},
        'RELEASEENAME': {'text': input_data.get("releasee-name", ""), 'font_size': 12, 'bold': True},
        'RELEASEEFATHER': {'text': input_data.get("releasees-fathers-name", ""), 'font_size': 12, 'bold': True},
        'RELEASEEADDRESS': {'text': input_data.get("releasees-address", ""), 'font_size': 12, 'bold': True},
        'PROPERTYSQFT': {'text': input_data.get("property-size", ""), 'font_size': 12, 'bold': True},
        'DOORNO': {'text': input_data.get("property-door-number", ""), 'font_size': 12, 'bold': True},
        'ROADNAME': {'text': input_data.get("property-location", ""), 'font_size': 12, 'bold': True},
        'SURVEYNO': {'text': input_data.get("survey-number", ""), 'font_size': 12, 'bold': True},
        'TALUKNAME': {'text': input_data.get("taluk-name", ""), 'font_size': 12, 'bold': True},
        'DISTRICTNAME': {'text': input_data.get("district-name", ""), 'font_size': 12, 'bold': True},
        'DOCNO': {'text': input_data.get("document-number", ""), 'font_size': 12, 'bold': True},
        'FILEPAGESTART': {'text': input_data.get("registration-file-page-start", ""), 'font_size': 12, 'bold': True},
        'FILEPAGEEND': {'text': input_data.get("registration-file-page-end", ""), 'font_size': 12, 'bold': True},
        'SUBREGISTRARPLACE': {'text': input_data.get("sub-registrar-office-location", ""), 'font_size': 12, 'bold': True},
        'RIGHTSQUAREFEET': {'text': input_data.get("rights-granted", ""), 'font_size': 12, 'bold': True},
    }

    # Specify the path to your template file
    template_file_path = r"D:\Legex\Legex\backend-auth\templates\Release Deed.docx"
    summary_template_file_path = r"D:\Legex\Legex\backend-auth\summary\Release Deed.txt"
    # Check if the summary template file exists
    if not os.path.exists(summary_template_file_path):
        print(json.dumps({"error": f"Summary template file '{summary_template_file_path}' not found."}))
        sys.exit(1)

    # Check if the template file exists
    if not os.path.exists(template_file_path):
        print(json.dumps({"error": f"Template file '{template_file_path}' not found."}))
        sys.exit(1)

    # Load the Word document template
    doc = docx.Document(template_file_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        replace_placeholders(paragraph, data)

    # Save the filled-in document
    output_file_path = r"D:\Legex\Legex\backend-auth\filled_documents\Filled_document_release_deed.docx"
    doc.save(output_file_path)

    # Load the summary template file
    summary = open(summary_template_file_path, 'r').read()
    filled_summary = replace_placeholders_in_text(summary, data)

    # Save the summary to a text file
    summary_output_file_path = r"D:\Legex\Legex\backend-auth\filled_summary\Filled_document_release_deed_summary.txt"
    with open(summary_output_file_path, 'w') as summary_file:
        summary_file.write(filled_summary)
        
    # Convert the Word document to PDF
    pdf_output_file_path = r"D:\Legex\Legex\backend-auth\filled_documents\Filled_document_release_deed.pdf"
    convert_to_pdf(output_file_path, pdf_output_file_path)
    
    #roadmap folder path
    roadmap_folder_path = r"D:\Legex\Legex\backend-auth\Roadmap\release deed"

    # Return the output file path
    print(json.dumps({
            "wordFilePath": output_file_path,
            "summaryFilePath": summary_output_file_path,
            "pdfFilePath":pdf_output_file_path,
            "roadmapFolderPath": roadmap_folder_path
        }))
import docx
import os
import sys
import json
from docx2pdf import convert

# Function to replace placeholders with data and apply formatting
def replace_text_and_format(run, text, font_size=12, bold=False, italic=False):
    run.text = text  # Replace the text
    run.font.size = docx.shared.Pt(font_size)  # Set font size
    run.bold = bold  # Set bold
    run.italic = italic  # Set italic

# Function to replace placeholders in a paragraph
def replace_placeholders(paragraph, data):
    for key, value_info in data.items():
        if key in paragraph.text:
            for run in paragraph.runs:
                if key in run.text:
                    replace_text_and_format(
                        run,
                        value_info['text'],
                        font_size=value_info['font_size'],
                        bold=value_info['bold'],
                    )
def replace_placeholders_in_text(text, data):
    for key, value_info in data.items():
        text = text.replace(key, value_info['text'])  # Replace the placeholder with its value
    return text

# Function to convert Word document to PDF
def convert_to_pdf(word_file_path, pdf_file_path):
    convert(word_file_path, pdf_file_path)

if __name__ == "__main__":
    # Read JSON input from stdin
    input_data = json.loads(sys.stdin.read())

    # Map input data to the placeholders in the template
    data = {
        'PLACE': {'text': input_data.get("place-of-agreement", ""), 'font_size': 12, 'bold': True},
        'DATE': {'text': input_data.get("agreement-day", ""), 'font_size': 12, 'bold': True},
        'YEAR': {'text': input_data.get("agreement-year", ""), 'font_size': 12, 'bold': True},
        'TRUSTORNAME': {'text': input_data.get("trustors-name", ""), 'font_size': 12, 'bold': True},
        'TRUSTORFATHER': {'text': input_data.get("trustors-fathers-name", ""), 'font_size': 12, 'bold': True},
        'TRUSTORADDRESS': {'text': input_data.get("trustors-address", ""), 'font_size': 12, 'bold': True},
        'TRUSTFUND': {'text': input_data.get("trust-fund-amount", ""), 'font_size': 12, 'bold': True},
        'TRUSTFUNDWORDS': {'text': input_data.get("trust-fund-amount-in-words", ""), 'font_size': 12, 'bold': True},
        'ADVANCEPAYMENT': {'text': input_data.get("advance-payment-amount", ""), 'font_size': 12, 'bold': True},
        'TRUSTADDRESS': {'text': input_data.get("trust-address", ""), 'font_size': 12, 'bold': True},
        'TRUSTAUTHORITIES': {'text': input_data.get("trust-authorities-names", ""), 'font_size': 12, 'bold': True},
        'TRUSTOBJECTIVES': {'text': input_data.get("trust-objectives", ""), 'font_size': 12, 'bold': True},
        'TRUSTFUNCTIONS': {'text': input_data.get("trust-functions", ""), 'font_size': 12, 'bold': True},
        'TRUSTOCCUPATIONS': {'text': input_data.get("trust-occupations", ""), 'font_size': 12, 'bold': True},
        'REMOVALCONDITIONS': {'text': input_data.get("conditions-of-trustee-removal", ""), 'font_size': 12, 'bold': True},
        'NEWTRUSTEES': {'text': input_data.get("appointment-of-new-trustees", ""), 'font_size': 12, 'bold': True},
        'TRUSTADMIN': {'text': input_data.get("trust-administration-details", ""), 'font_size': 12, 'bold': True},
        'MEETINGSCHEDULE': {'text': input_data.get("trust-meeting-schedule", ""), 'font_size': 12, 'bold': True},
        'RESOLUTIONS': {'text': input_data.get("trust-resolutions-decision-making", ""), 'font_size': 12, 'bold': True},
        'LEGALRIGHTS': {'text': input_data.get("trust-legal-rights", ""), 'font_size': 12, 'bold': True},
        'BANKDETAILS': {'text': input_data.get("trust-bank-account-details", ""), 'font_size': 12, 'bold': True},
        'INDEMNITY': {'text': input_data.get("trust-indemnity-conditions", ""), 'font_size': 12, 'bold': True},
        'ACTIVITIES': {'text': input_data.get("trust-activities-operations", ""), 'font_size': 12, 'bold': True},
        'DISSOLUTION': {'text': input_data.get("trust-dissolution-procedure", ""), 'font_size': 12, 'bold': True},
        'FUNDPROCEEDINGS': {'text': input_data.get("trust-fund distribution-proceedings", ""), 'font_size': 12, 'bold': True},
    }

    # Specify the path to your template file
    template_file_path = r"D:\Legex-Integration\Legex\backend-auth\templates\trust deed.docx"
    summary_template_file_path = r"D:\Legex-Integration\Legex\backend-auth\summary\trust deed.txt"

    # Check if the summary template file exists
    if not os.path.exists(summary_template_file_path):
        print(json.dumps({"error": f"Summary template file '{summary_template_file_path}' not found."}))
        sys.exit(1)

    # Check if the template file exists
    if not os.path.exists(template_file_path):
        print(json.dumps({"error": f"Template file '{template_file_path}' not found."}))
        sys.exit(1)

    # Load the Word document template
    doc = docx.Document(template_file_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        replace_placeholders(paragraph, data)

    # Save the filled-in document
    output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_trust_deed.docx"
    doc.save(output_file_path)

    # Load the summary template file
    summary = open(summary_template_file_path, 'r').read()
    filled_summary = replace_placeholders_in_text(summary, data)

    # Save the summary to a text file
    summary_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_summary\Filled_document_trust_deed.txt"
    with open(summary_output_file_path, 'w') as summary_file:
        summary_file.write(filled_summary)

    # Convert the Word document to PDF
    pdf_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_trust_deed.pdf"
    convert_to_pdf(output_file_path, pdf_output_file_path)
    
    #Roadmap folder path
    roadmap_folder_path = r"D:\Legex-Integration\Legex\backend-auth\Roadmap\Trust deed"
    # Return the output file path
    print(json.dumps({
            "wordFilePath": output_file_path,
            "summaryFilePath": summary_output_file_path,
            "pdfFilePath":pdf_output_file_path,
            "roadmapFolderPath": roadmap_folder_path
        }))
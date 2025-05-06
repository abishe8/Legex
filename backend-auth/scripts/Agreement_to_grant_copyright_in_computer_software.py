import docx
import os
import sys
import json
from docx2pdf import convert
from deep_translator import GoogleTranslator

# Function to replace placeholders with data and apply formatting
def replace_text_and_format(run, text, font_size=12, bold=False, italic=False):
    run.text = text  # Replace the text
    run.font.size = docx.shared.Pt(font_size)  # Set font size
    run.bold = bold  # Set bold
    run.italic = italic  # Set italic

# Function to replace placeholders in a paragraph
def replace_placeholders(paragraph, data):
    for key, value_info in data.items():
        if key in paragraph.text:
            for run in paragraph.runs:
                if key in run.text:
                    replace_text_and_format(
                        run,
                        value_info['text'],
                        font_size=value_info['font_size'],
                        bold=value_info['bold'],
                    )
def replace_placeholders_in_text(text, data):
    for key, value_info in data.items():
        text = text.replace(key, value_info['text'])  # Replace the placeholder with its value
    return text

# Function to convert Word document to PDF
def convert_to_pdf(word_file_path, pdf_file_path):
    convert(word_file_path, pdf_file_path)


if __name__ == "__main__":
    # Read JSON input from stdin
    input_data = json.loads(sys.stdin.read())

    # Map input data to the placeholders in the template
    data = {
        'PLACE': {'text': input_data.get("place", ""), 'font_size': 12, 'bold': True},
        'DATE': {'text': input_data.get("agreement-day", ""), 'font_size': 12, 'bold': True},
        'MONTH': {'text': input_data.get("agreement-month", ""), 'font_size': 12, 'bold': True},
        'YEAR': {'text': input_data.get("agreement-year", ""), 'font_size': 12, 'bold': True},
        'CONSULTANTNAME': {'text': input_data.get("consultant-name", ""), 'font_size': 12, 'bold': True},
        'PROGRAMMERNAME': {'text': input_data.get("programmer-name", ""), 'font_size': 12, 'bold': True},
        'CONSULTANTADDRESS': {'text': input_data.get("consultant-address", ""), 'font_size': 12, 'bold': True},
        'PROGRAMMERFATHER': {'text': input_data.get("programmer-father", ""), 'font_size': 12, 'bold': True},
        'PROGRAMMERADDRESS': {'text': input_data.get("programmer-address", ""), 'font_size': 12, 'bold': True},
        'PROGRAMMERPAYMENT': {'text': input_data.get("programmer-payment", ""), 'font_size': 12, 'bold': True},
        'DELIVERYYEARMONTH': {'text': input_data.get("delivery-year-month", ""), 'font_size': 12, 'bold': True},
        'DATEDELIVERY': {'text': input_data.get("delivery-day", ""), 'font_size': 12, 'bold': True},
        'FIRSTADVANCE': {'text': input_data.get("first-advance-payment", ""), 'font_size': 12, 'bold': True},
        'REMAININGPAYMENT': {'text': input_data.get("remaining-payment", ""), 'font_size': 12, 'bold': True},
    }

    # Specify the path to your template file
    template_file_path = r"D:\Legex-Integration\Legex\backend-auth\templates\copyright in computer software.docx"
    summary_template_file_path = r"D:\Legex-Integration\Legex\backend-auth\summary\copyright in computer software.txt"
   
    # Check if the template file exists
    if not os.path.exists(template_file_path):
        print(json.dumps({"error": f"Template file '{template_file_path}' not found."}))
        sys.exit(1)

    # Load the Word document template
    doc = docx.Document(template_file_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        replace_placeholders(paragraph, data)

    # Save the filled-in document
    output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_agreement_copyright.docx"
    doc.save(output_file_path)
    
    #Load the summary template file
    if not os.path.exists(summary_template_file_path):
        print(json.dumps({"error": f"Summary template file '{summary_template_file_path}' not found."}))
        sys.exit(1)
        
    summary = open(summary_template_file_path, 'r').read()
    filled_summary = replace_placeholders_in_text(summary, data)

    #Save the summary to a text file
    summary_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_summary\Filled_document_agreement_copyright.txt"
    with open(summary_output_file_path, 'w') as summary_file:
        summary_file.write(filled_summary)

    language_code_map = {
    "hindi": "hi",
    "tamil": "ta",
    "telugu": "te",
    "bengali": "bn",
    "english": "en"
    }
    selected_language = input_data.get("select-language", "english").lower()
    target_lang_code = language_code_map.get(selected_language, "en")

    translated_summary = filled_summary
    if target_lang_code != "en":
        try:
            translated_summary = GoogleTranslator(source='auto', target=target_lang_code).translate(filled_summary)
        except Exception as e:
            translated_summary = f"Translation failed: {str(e)}"

    # Save the translated summary
    translated_summary_path = rf"D:\Legex-Integration\Legex\backend-auth\filled_summary\Translated_document_agreement_copyright_{selected_language}.txt"
    with open(translated_summary_path, 'w', encoding='utf-8') as file:
        file.write(translated_summary)

    pdf_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_agreement_copyright.pdf"
    convert(output_file_path, pdf_output_file_path)

    #Roadmap folder path
    roadmap_folder_path = r"D:\Legex-Integration\Legex\backend-auth\Roadmap\Agreement to grant copyrights to computers"
    print(json.dumps({
            "wordFilePath": output_file_path,
            "summaryFilePath": summary_output_file_path,
            "pdfFilePath": pdf_output_file_path,
            "roadmapFolderPath": roadmap_folder_path,
            "translatedSummaryFilePath": translated_summary_path,
        }))
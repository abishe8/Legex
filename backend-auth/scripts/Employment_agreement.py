import docx
import os
import sys
import json
from docx2pdf import convert

# Function to replace placeholders with data and apply formatting
def replace_text_and_format(run, text, font_size=12, bold=False, italic=False):
    run.text = text  # Replace the text
    run.font.size = docx.shared.Pt(font_size)  # Set font size
    run.bold = bold  # Set bold
    run.italic = italic  # Set italic

# Function to replace placeholders in a paragraph
def replace_placeholders(paragraph, data):
    for key, value_info in data.items():
        if key in paragraph.text:
            for run in paragraph.runs:
                if key in run.text:
                    replace_text_and_format(
                        run,
                        value_info['text'],
                        font_size=value_info['font_size'],
                        bold=value_info['bold'],
                    )
def replace_placeholders_in_text(text, data):
    for key, value_info in data.items():
        text = text.replace(key, value_info['text'])  # Replace the placeholder with its value
    return text

# Function to convert Word document to PDF
def convert_to_pdf(word_file_path, pdf_file_path):
    convert(word_file_path, pdf_file_path)

if __name__ == "__main__":
    # Read JSON input from stdin
    input_data = json.loads(sys.stdin.read())

    # Map input data to the placeholders in the template
    data = {
        'DATE': {'text': input_data.get("agreement-date", ""), 'font_size': 12, 'bold': True},
        'MONTH': {'text': input_data.get("agreement-month", ""), 'font_size': 12, 'bold': True},
        'YEAR': {'text': input_data.get("agreement-year", ""), 'font_size': 12, 'bold': True},
        'COMPANYNAME': {'text': input_data.get("company-name", ""), 'font_size': 12, 'bold': True},
        'COMPANYADDRESS': {'text': input_data.get("company-address", ""), 'font_size': 12, 'bold': True},
        'EMPLOYEENAME': {'text': input_data.get("employee-name", ""), 'font_size': 12, 'bold': True},
        'DEPENDENTNAME': {'text': input_data.get("dependant-name", ""), 'font_size': 12, 'bold': True},
        'AGE': {'text': input_data.get("employee-age", ""), 'font_size': 12, 'bold': True},
        'ADDRESS': {'text': input_data.get("employee-address", ""), 'font_size': 12, 'bold': True},
        'POST': {'text': input_data.get("job-position", ""), 'font_size': 12, 'bold': True},
        'INITIALDAYS': {'text': input_data.get("initial-work-period", ""), 'font_size': 12, 'bold': True},
        'RESPONSIBILITY': {'text': input_data.get("job-responsibilities", ""), 'font_size': 12, 'bold': True},
        'SAL': {'text': input_data.get("salary-amount", ""), 'font_size': 12, 'bold': True},
        'CASUALLEAVE': {'text': input_data.get("causal-leave-days", ""), 'font_size': 12, 'bold': True},
        'CWORD': {'text': input_data.get("causal-leave-in-words", ""), 'font_size': 12, 'bold': True},
        'SICKLEAVE': {'text': input_data.get("sick-leave-days", ""), 'font_size': 12, 'bold': True},
        'SWORD': {'text': input_data.get("sick-leave-in-words", ""), 'font_size': 12, 'bold': True},
        'PUBHOLI': {'text': input_data.get("public-holidays-allowed", ""), 'font_size': 12, 'bold': True},
        'PWORD': {'text': input_data.get("public-holidays-in-words", ""), 'font_size': 12, 'bold': True},
        'NOTICEPERIOD': {'text': input_data.get("notice-period-duration", ""), 'font_size': 12, 'bold': True},
        'NOTICEWORD': {'text': input_data.get("notice-period-in-words", ""), 'font_size': 12, 'bold': True},
        'EADVANCE': {'text': input_data.get("employee-advance-payment", ""), 'font_size': 12, 'bold': True},
        'EADWORD': {'text': input_data.get("employee-advance-payment-words", ""), 'font_size': 12, 'bold': True},
    }

    # Specify the path to your template file
    template_file_path = r"D:\Legex-Integration\Legex\backend-auth\templates\Employment agreement.docx"
    summary_template_file_path = r"D:\Legex-Integration\Legex\backend-auth\summary\Employment agreement.txt"

    # Check if the summary template file exists
    if not os.path.exists(summary_template_file_path):
        print(json.dumps({"error": f"Summary template file '{summary_template_file_path}' not found."}))
        sys.exit(1)

    # Check if the template file exists
    if not os.path.exists(template_file_path):
        print(json.dumps({"error": f"Template file '{template_file_path}' not found."}))
        sys.exit(1)

    # Load the Word document template
    doc = docx.Document(template_file_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        replace_placeholders(paragraph, data)

    # Save the filled-in document
    output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_employment_agreement.docx"
    doc.save(output_file_path)

    # Load the summary template file
    summary = open(summary_template_file_path, 'r', encoding='utf-8').read()
    filled_summary = replace_placeholders_in_text(summary, data)

    # Save the summary to a text file
    summary_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_summary\Filled_document_employment_agreement.txt"
    with open(summary_output_file_path, 'w') as summary_file:
        summary_file.write(filled_summary)
        
    # Convert the Word document to PDF
    pdf_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_employment_agreement.pdf"
    convert_to_pdf(output_file_path, pdf_output_file_path)
    
    #Roadmap folder path
    roadmap_folder_path = r"D:\Legex-Integration\Legex\backend-auth\Roadmap\Employment agreement"
    # Return the output file path
    print(json.dumps({
            "wordFilePath": output_file_path,
            "summaryFilePath": summary_output_file_path,
            "pdfFilePath":pdf_output_file_path,
            "roadmapFolderPath": roadmap_folder_path
        }))
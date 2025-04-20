import docx
import os
import sys
import json
from docx2pdf import convert

# Function to replace placeholders with data and apply formatting
def replace_text_and_format(run, text, font_size=12, bold=False, italic=False):
    run.text = text  # Replace the text
    run.font.size = docx.shared.Pt(font_size)  # Set font size
    run.bold = bold  # Set bold
    run.italic = italic  # Set italic

# Function to replace placeholders in a paragraph
def replace_placeholders(paragraph, data):
    for key, value_info in data.items():
        if key in paragraph.text:
            for run in paragraph.runs:
                if key in run.text:
                    replace_text_and_format(
                        run,
                        value_info['text'],
                        font_size=value_info['font_size'],
                        bold=value_info['bold'],
                    )
def replace_placeholders_in_text(text, data):
    for key, value_info in data.items():
        text = text.replace(key, value_info['text'])  # Replace the placeholder with its value
    return text

# Function to convert Word document to PDF
def convert_to_pdf(word_file_path, pdf_file_path):
    convert(word_file_path, pdf_file_path)

if __name__ == "__main__":
    # Read JSON input from stdin
    input_data = json.loads(sys.stdin.read())

    # Map input data to the placeholders in the template
    data = {
        'DATE': {'text': input_data.get("agreement-day", ""), 'font_size': 12, 'bold': True},
        'MONTH': {'text': input_data.get("agreement-month", ""), 'font_size': 12, 'bold': True},
        'YEAR': {'text': input_data.get("agreement-year", ""), 'font_size': 12, 'bold': True},
        'BORROWER': {'text': input_data.get("borrowers-name", ""), 'font_size': 12, 'bold': True},
        'LOCATIONB': {'text': input_data.get("borrowers-location", ""), 'font_size': 12, 'bold': True},
        'LENDER': {'text': input_data.get("lenders-name", ""), 'font_size': 12, 'bold': True},
        'ADDRESSL': {'text': input_data.get("lenders-address", ""), 'font_size': 12, 'bold': True},
        'LENDAMT': {'text': input_data.get("loan-amount", ""), 'font_size': 12, 'bold': True},
        'DA': {'text': input_data.get("loan-agreement-day", ""), 'font_size': 12, 'bold': True},
        'MON': {'text': input_data.get("loan-agreement-month", ""), 'font_size': 12, 'bold': True},
        'YR': {'text': input_data.get("loan-agreement-year", ""), 'font_size': 12, 'bold': True},
        'INTERESTRATE': {'text': input_data.get("interest-rate", ""), 'font_size': 12, 'bold': True},
        'FIXEDDAY': {'text': input_data.get("fixed-repayment-day", ""), 'font_size': 12, 'bold': True},
    }

    # Specify the path to your template file
    template_file_path = r"D:\Legex-Integration\Legex\backend-auth\templates\Deed of Simple Mortgage.docx"
    summary_template_file_path = r"D:\Legex-Integration\Legex\backend-auth\summary\deed_of_simple_mortgage.txt"
    # Check if the template file exists
    if not os.path.exists(template_file_path):
        print(json.dumps({"error": f"Template file '{template_file_path}' not found."}))
        sys.exit(1)

    # Load the Word document template
    doc = docx.Document(template_file_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        replace_placeholders(paragraph, data)

    # Save the filled-in document
    output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_deed_of_simple_mortgage.docx"
    doc.save(output_file_path)

    #check if the summary template file exists
    if not os.path.exists(summary_template_file_path):
        print(json.dumps({"error": f"Summary template file '{summary_template_file_path}' not found."}))
        sys.exit(1)

    summary = open(summary_template_file_path, 'r').read()
    filled_summary = replace_placeholders_in_text(summary, data)

    # Save the summary to a text file
    summary_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_summary\Filled_deed_of_simple_mortgage_summary.txt"
    with open(summary_output_file_path, 'w') as summary_file:
        summary_file.write(filled_summary)
    
    #Convert the Word document to PDF
    pdf_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_deed_of_simple_mortgage.pdf"
    convert_to_pdf(output_file_path, pdf_output_file_path)
    
    #Roadmap Path 
    roadmap_folder_path = r"D:\Legex-Integration\Legex\backend-auth\Roadmap\Deed on Simple Mortgage"
    # Return the output file path
    print(json.dumps({
            "wordFilePath": output_file_path,
            "summaryFilePath": summary_output_file_path,
            "pdfFilePath":pdf_output_file_path,
            "roadmapFolderPath": roadmap_folder_path
        }), flush=True)
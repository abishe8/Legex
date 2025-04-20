import docx
import os
import sys
import json
from docx2pdf import convert
from deep_translator import GoogleTranslator

# Function to replace placeholders with data and apply formatting
def replace_text_and_format(run, text, font_size=12, bold=False, italic=False):
    run.text = text  # Replace the text
    run.font.size = docx.shared.Pt(font_size)  # Set font size
    run.bold = bold  # Set bold
    run.italic = italic  # Set italic

# Function to replace placeholders in a paragraph
def replace_placeholders(paragraph, data):
    for key, value_info in data.items():
        if key in paragraph.text:
            for run in paragraph.runs:
                if key in run.text:
                    replace_text_and_format(
                        run,
                        value_info['text'],
                        font_size=value_info['font_size'],
                        bold=value_info['bold'],
                    )
def replace_placeholders_in_text(text, data):
    for key, value_info in data.items():
        text = text.replace(key, value_info['text'])  # Replace the placeholder with its value
    return text

# Function to convert Word document to PDF
def convert_to_pdf(word_file_path, pdf_file_path):
    convert(word_file_path, pdf_file_path)

if __name__ == "__main__":
    # Read JSON input from stdin
    input_data = json.loads(sys.stdin.read())

    # Map input data to the placeholders in the template
    data = {
        'PRINCIPALNAME': {'text': input_data.get("principal-name", ""), 'font_size': 12, 'bold': True},
        'PRINCIPALAGE': {'text': input_data.get("principal-age", ""), 'font_size': 12, 'bold': True},
        'PRINCIPALADDRESS': {'text': input_data.get("principal-address", ""), 'font_size': 12, 'bold': True},
        'NOMINEENAME': {'text': input_data.get("nominee-name", ""), 'font_size': 12, 'bold': True},
        'NOMINEEFATHER': {'text': input_data.get("nominee-fathers-name", ""), 'font_size': 12, 'bold': True},
        'NOMINEEAGE': {'text': input_data.get("nominee-age", ""), 'font_size': 12, 'bold': True},
        'NOMINEEADDRESS': {'text': input_data.get("nominee-address", ""), 'font_size': 12, 'bold': True},
        'BANKNAME': {'text': input_data.get("bank-name", ""), 'font_size': 12, 'bold': True},
        'DAY': {'text': input_data.get("day-of-agreement", ""), 'font_size': 12, 'bold': True},
        'YEAR': {'text': input_data.get("year-of-agreement", ""), 'font_size': 12, 'bold': True},
        'PLACE': {'text': input_data.get("place-of-agreement", ""), 'font_size': 12, 'bold': True},
    }

    # Specify the path to your template file
    template_file_path = r"D:\Legex-Integration\Legex\backend-auth\templates\Power of attorney.docx"
    summary_template_file_path = r"D:\Legex-Integration\Legex\backend-auth\summary\Power of attorney.txt"

    # Check if the summary template file exists
    if not os.path.exists(summary_template_file_path):
        print(json.dumps({"error": f"Summary template file '{summary_template_file_path}' not found."}))
        sys.exit(1)

    # Check if the template file exists
    if not os.path.exists(template_file_path):
        print(json.dumps({"error": f"Template file '{template_file_path}' not found."}))
        sys.exit(1)

    # Load the Word document template
    doc = docx.Document(template_file_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        replace_placeholders(paragraph, data)

    # Save the filled-in document
    output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_power_of_attorney.docx"
    doc.save(output_file_path)

    # Load the summary template file
    summary = open(summary_template_file_path, 'r').read()
    filled_summary = replace_placeholders_in_text(summary, data)

    summary_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_summary\Filled_document_power_of_attorney_summary.txt"
    with open(summary_output_file_path, 'w') as summary_file:
        summary_file.write(filled_summary)

        language_code_map = {
    "hindi": "hi",
    "tamil": "ta",
    "telugu": "te",
    "bengali": "bn",
    "english": "en"
    }
    selected_language = input_data.get("select-language", "english").lower()
    target_lang_code = language_code_map.get(selected_language, "en")

    translated_summary = filled_summary
    if target_lang_code != "en":
        try:
            translated_summary = GoogleTranslator(source='auto', target=target_lang_code).translate(filled_summary)
        except Exception as e:
            translated_summary = f"Translation failed: {str(e)}"

    # Save the translated summary
    translated_summary_path = rf"D:\Legex-Integration\Legex\backend-auth\filled_summary\Translated_power_of_attorney_{selected_language}.txt"
    with open(translated_summary_path, 'w', encoding='utf-8') as file:
        file.write(translated_summary)
    
    # Convert the Word document to PDF
    pdf_output_file_path = r"D:\Legex-Integration\Legex\backend-auth\filled_documents\Filled_document_power_of_attorney.pdf"
    convert_to_pdf(output_file_path, pdf_output_file_path)
    
    #roadmap folder path
    roadmap_folder_path = r"D:\Legex-Integration\Legex\backend-auth\Roadmap\Power of Attorney"
    # Return the output file path
    print(json.dumps({
            "wordFilePath": output_file_path,
            "summaryFilePath": summary_output_file_path,
            "pdfFilePath":pdf_output_file_path,
            "roadmapFolderPath": roadmap_folder_path,
            "translatedSummaryFilePath": translated_summary_path,
        }))